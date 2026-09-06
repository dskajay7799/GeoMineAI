import os
from pathlib import Path
from typing import Optional

import pandas as pd
import pytesseract
from docx import Document as DocxDocument
from pdf2image import convert_from_path
from PIL import Image
from pypdf import PdfReader

# ============================================================
# TESSERACT CONFIGURATION
# ============================================================

TESSERACT_PATH = os.environ.get(
    "TESSERACT_PATH",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
)

if os.path.exists(TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# POPPLER CONFIGURATION
# ============================================================

POPPLER_PATH: Optional[str] = os.environ.get(
    "POPPLER_PATH",
    os.path.join(
        os.path.expanduser("~"),
        "Downloads",
        "poppler-26.02.0",
        "Library",
        "bin",
    ),
)

if not os.path.exists(POPPLER_PATH):
    POPPLER_PATH = None  # let pdf2image fall back to system PATH


# ============================================================
# DIGITAL PDF
# ============================================================

def process_pdf(file_path: Path) -> dict:
    reader = PdfReader(file_path)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page": page_number, "text": text})

    full_text = "\n\n".join(page["text"] for page in pages)

    return {
        "document_type": "PDF",
        "text": full_text,
        "pages": pages,
        "page_count": len(pages),
        "ocr_used": False,
        "extraction_method": "PDF_TEXT",
    }


# ============================================================
# SCANNED PDF
# ============================================================

def process_scanned_pdf(file_path: Path) -> dict:
    images = convert_from_path(
        file_path,
        dpi=200,
        poppler_path=POPPLER_PATH,
    )

    extracted_pages = []

    for page_number, page_image in enumerate(images, start=1):
        text = pytesseract.image_to_string(page_image)
        extracted_pages.append({"page": page_number, "text": text})

    full_text = "\n\n".join(page["text"] for page in extracted_pages)

    return {
        "document_type": "SCANNED_PDF",
        "text": full_text,
        "pages": extracted_pages,
        "page_count": len(extracted_pages),
        "ocr_used": True,
        "extraction_method": "OCR",
    }


# ============================================================
# WORD
# ============================================================

def process_docx(file_path: Path) -> dict:
    document = DocxDocument(file_path)

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables = [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]

    full_text = "\n".join(paragraphs)

    return {
        "document_type": "DOCX",
        "text": full_text,
        "paragraph_count": len(paragraphs),
        "tables": tables,
        "pages": [{"page": 1, "text": full_text}],
        "page_count": 1,
        "ocr_used": False,
        "extraction_method": "DOCX_TEXT",
    }


# ============================================================
# EXCEL
# ============================================================

def _dataframe_to_text(sheet_name: str, dataframe: pd.DataFrame) -> str:
    """
    Convert a spreadsheet's rows into readable text so it can be
    searched and read by the AI Assistant / word cloud / reports,
    the same way PDF/DOCX text is used.
    """
    lines = [f"Sheet: {sheet_name}"]
    lines.append("Columns: " + ", ".join(str(c) for c in dataframe.columns))
    lines.append("")

    for _, row in dataframe.iterrows():
        row_parts = [
            f"{column}: {row[column]}"
            for column in dataframe.columns
            if str(row[column]).strip() not in ("", "nan")
        ]
        if row_parts:
            lines.append(" | ".join(row_parts))

    return "\n".join(lines)


def process_spreadsheet(file_path: Path) -> dict:
    excel_file = pd.ExcelFile(file_path)
    sheets = {}
    pages = []

    for page_number, sheet_name in enumerate(excel_file.sheet_names, start=1):
        dataframe = pd.read_excel(file_path, sheet_name=sheet_name)
        dataframe = dataframe.fillna("")

        sheets[sheet_name] = dataframe.to_dict(orient="records")

        sheet_text = _dataframe_to_text(sheet_name, dataframe)
        pages.append({"page": page_number, "text": sheet_text})

    full_text = "\n\n".join(page["text"] for page in pages)

    return {
        "document_type": "SPREADSHEET",
        "text": full_text,
        "sheets": sheets,
        "sheet_names": excel_file.sheet_names,
        "pages": pages,
        "page_count": len(pages),
        "ocr_used": False,
        "extraction_method": "SPREADSHEET",
    }


# ============================================================
# CSV
# ============================================================

def process_csv(file_path: Path) -> dict:
    dataframe = pd.read_csv(file_path)
    dataframe = dataframe.fillna("")

    full_text = _dataframe_to_text(file_path.stem, dataframe)

    return {
        "document_type": "CSV",
        "text": full_text,
        "records": dataframe.to_dict(orient="records"),
        "columns": list(dataframe.columns),
        "row_count": len(dataframe),
        "pages": [{"page": 1, "text": full_text}],
        "page_count": 1,
        "ocr_used": False,
        "extraction_method": "CSV",
    }


# ============================================================
# IMAGE OCR
# ============================================================

def process_image(file_path: Path) -> dict:
    image = Image.open(file_path)
    text = pytesseract.image_to_string(image)

    return {
        "document_type": "IMAGE",
        "text": text,
        "pages": [{"page": 1, "text": text}],
        "page_count": 1,
        "ocr_used": True,
        "extraction_method": "OCR",
    }


# ============================================================
# MAIN PROCESSOR
# ============================================================

SPREADSHEET_EXTENSIONS = {".xlsx", ".xls"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def process_document(file_path) -> dict:
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        result = process_pdf(file_path)
        if result["text"].strip():
            return result
        return process_scanned_pdf(file_path)

    if extension == ".docx":
        return process_docx(file_path)

    if extension in SPREADSHEET_EXTENSIONS:
        return process_spreadsheet(file_path)

    if extension == ".csv":
        return process_csv(file_path)

    if extension in IMAGE_EXTENSIONS:
        return process_image(file_path)

    return {
        "document_type": "UNSUPPORTED",
        "text": "",
        "pages": [],
        "page_count": 0,
        "ocr_used": False,
        "extraction_method": "UNKNOWN",
    }